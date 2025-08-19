import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

# Read the existing CSV file
df = pd.read_csv('toolspy-active-tools-schedule.csv')

# Create Excel workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Toolspy Active Tools'

# Add data to worksheet
for r in dataframe_to_rows(df, index=False, header=True):
    ws.append(r)

# Style the header row
header_font = Font(bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')

for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')

# Add borders and alignment to all cells
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
    for cell in row:
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

# Auto-adjust column widths
for column in ws.columns:
    max_length = 0
    column_letter = column[0].column_letter
    
    for cell in column:
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(str(cell.value))
        except:
            pass
    
    adjusted_width = min(max_length + 2, 50)
    ws.column_dimensions[column_letter].width = adjusted_width

# Set header row height
ws.row_dimensions[1].height = 30

# Save the file
wb.save('toolspy-active-tools-inventory.xlsx')
print('Excel file created successfully!')
print(f'Total tools in inventory: {len(df)}')
print(f'Active tools: {len(df[df["Status"] == "Active"])}')
print(f'Featured tools: {len(df[df["Featured"] == "Yes"])}')
print(f'Categories: {df["Category"].nunique()}')
